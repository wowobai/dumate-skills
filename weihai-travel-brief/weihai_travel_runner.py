#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
weihai_travel_runner.py V1 - 威海旅游新闻云端统一生成脚本
学习 hot_news_runner（HOT_NEWS V17+）架构：
  3合1: PPT + 公众号HTML + 公众号草稿(Supabase Edge Function直调, 单篇模式)
  - 自动依赖安装（ensure_dependencies）
  - 配图压缩 + 一致性校验
  - 蓝色系配色（ACC=#1A56DB），与 weihai_travel_news SKILL V1 一致
  - 草稿标题固定：{date} 威海旅游新闻（main_title 字段，严禁混用其他主题标题）
  - 凭据不硬编码：从环境变量读取，缺失时回退从同目录 hot_news_runner.py 提取

用法: python -X utf8 weihai_travel_runner.py news_data.json
"""

import json, os, sys, ssl, re, subprocess
from pathlib import Path

# ======================== Auto Install Dependencies ========================
def ensure_dependencies():
    """自动检测并安装缺失的Python依赖包，消除沙箱首次运行ModuleNotFoundError硬故障。"""
    required = {
        'requests': 'requests',
        'PIL': 'Pillow',
        'pptx': 'python-pptx',
    }
    missing = []
    for mod, pkg in required.items():
        try:
            __import__(mod)
        except ImportError:
            missing.append(pkg)
    if missing:
        print(f"--- 安装缺失依赖: {', '.join(missing)} ---")
        subprocess.check_call(
            [sys.executable, '-m', 'pip', 'install', '--quiet'] + missing,
            stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL
        )
        print("  依赖安装完成")

ensure_dependencies()

import requests
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_AUTO_SIZE

# ======================== Config ========================
BG = RGBColor(0xF5, 0xF5, 0xFA)
TITLE_C = RGBColor(0x1A, 0x56, 0xDB)
BODY_C = RGBColor(0x33, 0x33, 0x33)
ACCENT_C = RGBColor(0x3B, 0x82, 0xF6)
GRAY_C = RGBColor(0x99, 0x99, 0x99)
SAFE_GAP = Inches(0.12)
CHARS_PER_INCH = 4.2

MAIN_TITLE = "威海旅游新闻"


# ======================== Credentials ========================
def _read_const_from_file(path, name):
    """从指定 py 文件源码中提取顶层字符串常量（避免在新文件中硬编码凭据）。"""
    try:
        src = open(path, encoding="utf-8").read()
        m = re.search(rf'^{name}\s*=\s*["\']([^"\']+)["\']', src, re.M)
        return m.group(1) if m else ""
    except Exception:
        return ""


def load_credentials():
    """加载微信/秒哒凭据：优先环境变量，其次同目录 hot_news_runner.py（既有凭据载体）。"""
    creds = {
        "WX_APP_ID": os.environ.get("WX_APP_ID", ""),
        "WX_APP_SECRET": os.environ.get("WX_APP_SECRET", ""),
        "SUPABASE_URL": os.environ.get("SUPABASE_URL", ""),
        "SUPABASE_KEY": os.environ.get("SUPABASE_KEY", ""),
    }
    base = os.path.dirname(os.path.abspath(__file__))
    hot_src = os.path.join(base, "hot_news_runner.py")
    if os.path.exists(hot_src):
        for name in creds:
            if not creds[name]:
                creds[name] = _read_const_from_file(hot_src, name)
    if not creds["SUPABASE_KEY"]:
        print("  [WARN] SUPABASE_KEY 未配置（环境变量/同目录 hot_news_runner.py 均无），草稿环节将跳过")
    return creds


# ======================== Helpers ========================
def num_to_cn(n):
    d = "零一二三四五六七八九"
    if n < 10:
        return d[n]
    if n < 20:
        return "十" + (d[n % 10] if n % 10 else "")
    if n < 100:
        r = d[n // 10] + "十"
        return r + (d[n % 10] if n % 10 else "")
    return str(n)


def date_to_cn(date_str):
    """20260916 -> 二零二六年九月十六日"""
    y, m, d = int(date_str[:4]), int(date_str[4:6]), int(date_str[6:8])
    y_cn = "".join(num_to_cn(int(c)) for c in str(y))
    return f"{y_cn}年{num_to_cn(m)}月{num_to_cn(d)}日"


def truncate(text, limit=200):
    return text[:limit - 3] + "..." if len(text) > limit else text


_LANCZOS = getattr(Image, "LANCZOS", None) or getattr(getattr(Image, "Resampling", Image), "LANCZOS", 1)


def compress_image(path):
    img = Image.open(path)
    if img.mode in ('P', 'LA', 'RGBA', 'L', 'ARGB'):
        img = img.convert('RGB')
    if img.width > 1200:
        ratio = 1200 / img.width
        img = img.resize((1200, int(img.height * ratio)), _LANCZOS)
    img.save(path, "JPEG", quality=85)


# ======================== PPT ========================
def gen_ppt(base, data, news):
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)
    date_cn = data.get("date_chinese", date_to_cn(data["date"]))

    def add_bg(slide):
        bg = slide.background.fill
        bg.solid()
        bg.fore_color.rgb = BG

    def add_text(slide, x, y, w, h, text, size, color, bold=False, align=PP_ALIGN.LEFT, wrap=True, spacing=None):
        tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = wrap
        tf.auto_size = MSO_AUTO_SIZE.NONE
        p = tf.paragraphs[0]
        p.text = text
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.font.bold = bold
        p.alignment = align
        if spacing:
            p.line_spacing = spacing
        return p

    # Slide 1: Cover
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(s)
    add_text(s, 1, 2.3, 8, 1.5, "威海旅游新闻热点TOP5", 44, TITLE_C, True, PP_ALIGN.CENTER)
    add_text(s, 1, 3.4, 8, 0.8, date_cn, 24, ACCENT_C, False, PP_ALIGN.CENTER)

    # Slide 2: TOC
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(s)
    add_text(s, 0.5, 0.3, 9, 0.8, "今日导览", 32, TITLE_C, True)
    y = 1.4
    for i, item in enumerate(news):
        tag = item.get("tag", "")
        title = item["title"]
        add_text(s, 1, y, 8, 0.9, f"{i + 1}. [{tag}] {title}", 18, BODY_C)
        y += 0.95

    # Slides 3-7: Detail
    for item in news:
        s = prs.slides.add_slide(prs.slide_layouts[6])
        add_bg(s)
        img_path = os.path.join(base, item.get("image_path", f"images/news_{item['id']}.jpg"))
        if os.path.exists(img_path):
            try:
                img = Image.open(img_path)
                iw, ih = img.size
                ratio = min(4.5 / iw * 72, 4.5 / ih * 72) / 72
                s.shapes.add_picture(img_path, Inches(0.5), Inches(1.5),
                                     Inches(iw * ratio / 72), Inches(ih * ratio / 72))
            except Exception:
                pass

        body = item.get("body", item.get("summary", ""))
        analysis = item.get("analysis", item.get("impact", ""))
        add_text(s, 5, 0.5, 4.5, 1.5, f"[{item.get('tag','')}] {item['title']}", 20, TITLE_C, True)
        add_text(s, 5, 2.2, 4.5, 2.5, truncate(body), 14, BODY_C, spacing=1.35)
        add_text(s, 5, 5.0, 4.5, 1.5, "【看点】" + truncate(analysis, 150), 13, ACCENT_C, spacing=1.35)
        add_text(s, 5, 6.6, 4.5, 0.5, item["source"], 10, GRAY_C)

    # Slide 8: End
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(s)
    add_text(s, 1, 3, 8, 1.5, date_cn, 28, TITLE_C, False, PP_ALIGN.CENTER)
    add_text(s, 1, 4, 8, 0.8, "本期完", 20, ACCENT_C, False, PP_ALIGN.CENTER)

    path = os.path.join(base, "威海旅游新闻热点TOP5.pptx")
    prs.save(path)
    return path


# ======================== HTML ========================
def gen_html(base, data, news):
    date_cn = data.get("date_chinese", date_to_cn(data["date"]))
    dd = data["date_display"]

    css = ("body{font-family:-apple-system,'Microsoft YaHei',sans-serif;margin:0;padding:20px;color:#333}"
           ".h{text-align:center;padding:30px 0;background:linear-gradient(135deg,#1a56db,#3b82f6);color:#fff;border-radius:8px 8px 0 0}"
           ".h h1{font-size:28px;margin:0}.h .d{font-size:16px;margin-top:10px;opacity:.9}"
           ".ni{padding:25px 20px;border-bottom:1px solid #eee}"
           ".nt{font-size:20px;font-weight:bold;color:#1a56db;margin-bottom:12px}"
           ".ntg{font-size:12px;color:#fff;background:#1a56db;display:inline-block;padding:2px 10px;border-radius:4px;margin-bottom:6px}"
           ".nb{font-size:15px;line-height:1.8;margin-bottom:10px}"
           ".na{font-size:14px;color:#3b82f6;line-height:1.7;padding:8px 12px;background:#f0f7ff;border-radius:4px;margin-bottom:8px}"
           ".ns{font-size:12px;color:#999;font-style:italic}"
           ".ni img{width:100%;border-radius:6px;margin:12px 0}"
           ".ft{text-align:center;padding:20px;color:#999;font-size:13px}")

    parts = [f'<!DOCTYPE html><html lang="zh-CN"><head><meta charset="UTF-8">',
             f'<meta name="viewport" content="width=device-width,initial-scale=1">',
             f'<title>{dd} 威海旅游新闻</title><style>{css}</style></head><body>',
             f'<div style="max-width:677px;margin:0 auto">',
             f'<div class="h"><h1>威海旅游新闻热点TOP5</h1><div class="d">{date_cn}</div></div>']

    for item in news:
        parts.append('<div class="ni">')
        tag = item.get("tag", "")
        if tag:
            parts.append(f'<span class="ntg">{tag}</span>')
        parts.append(f'<div class="nt">{item["title"]}</div>')
        img_path = item.get("image_path", f"images/news_{item['id']}.jpg")
        parts.append(f'<img src="{img_path}" alt="{item["title"]}">')
        body = item.get("body", item.get("summary", ""))
        analysis = item.get("analysis", item.get("impact", ""))
        parts.append(f'<div class="nb">{truncate(body, 500)}</div>')
        parts.append(f'<div class="na">【旅游看点】{truncate(analysis, 300)}</div>')
        parts.append(f'<div class="ns">{item["source"]}</div>')
        parts.append('</div>')

    parts.append(f'<div class="ft"><p>{date_cn}</p><p>本期完</p></div>')
    parts.append('</div></body></html>')

    path = os.path.join(base, "wechat_article.html")
    with open(path, "w", encoding="utf-8") as f:
        f.write("\n".join(parts))
    return path


# ======================== WeChat Draft (Supabase Edge Function) ========================
def gen_wechat_draft(base, data, news, creds):
    """通过Supabase Edge Function创建公众号草稿（单篇模式，标题固定"威海旅游新闻"）。"""
    if not creds.get("SUPABASE_KEY"):
        print("  [SKIP] 未获取到 SUPABASE_KEY，跳过草稿创建")
        return None

    # 构造兼容 news_data（保留原字段 + 添加 body/analysis/image_query 别名）
    mapped_news = []
    for n in news:
        item = dict(n)
        item.setdefault("body", n.get("summary", ""))
        item.setdefault("analysis", n.get("impact", ""))
        item.setdefault("image_query", n.get("img_keyword", n.get("title", "")[:20]))
        mapped_news.append(item)

    payload_data = {
        "date": data["date"],
        "date_display": data["date_display"],
        "main_title": MAIN_TITLE,
        "news": mapped_news,
    }
    news_json = json.dumps(payload_data, ensure_ascii=False)

    img_dir = os.path.join(base, "images")
    files = []
    for i in range(1, len(news) + 1):
        img_path = os.path.join(img_dir, f"news_{i}.jpg")
        if os.path.exists(img_path):
            files.append(("images", (f"news_{i}.jpg", open(img_path, "rb"), "image/jpeg")))

    supabase_url = creds.get("SUPABASE_URL", "").rstrip("/")
    miaoda_url = f"{supabase_url}/functions/v1/create-draft"

    try:
        headers = {
            "Authorization": f"Bearer {creds['SUPABASE_KEY']}",
            "apikey": creds["SUPABASE_KEY"],
        }
        resp = requests.post(miaoda_url, data={"news_data": news_json}, files=files,
                             headers=headers, verify=False, timeout=180)
        result = resp.json()
        if result.get("success"):
            print(f"  [OK] 草稿创建成功! Draft ID: {result.get('draft_id')}")
            title = result.get("title", "")
            if title:
                print(f"  [OK] 草稿标题: {title}")
            else:
                print("  [WARN] 秒哒未回传 title 字段（旧版函数？）")
            verify = {
                "draft_media_id": result.get("draft_id"),
                "draft_title": title,
                "channel": "miaoda",
                "wx_img_urls_count": sum(1 for u in (result.get("content_images") or []) if u),
                "thumb_media_id": result.get("cover_media_id", ""),
            }
            verify_path = os.path.join(base, "draft_verify.json")
            with open(verify_path, "w", encoding="utf-8") as f:
                json.dump(verify, f, ensure_ascii=False, indent=2)
            return result
        else:
            print(f"  [ERROR] 草稿创建失败: {result.get('error', result)}")
            return None
    except Exception as e:
        print(f"  [ERROR] 秒哒代理请求异常: {e}")
        return None
    finally:
        for _, ft in files:
            ft[1].close()


# ======================== Verify ========================
def verify_all(base, data, news, ppt_path, html_path, draft_result=None):
    print("\n=== 验证产出 ===")
    ok = True

    if os.path.exists(ppt_path):
        size = os.path.getsize(ppt_path)
        prs = Presentation(ppt_path)
        pics = sum(1 for s in prs.slides for sh in s.shapes if sh.shape_type == 13)
        status = "OK" if size > 102400 and len(prs.slides) >= 8 and pics >= 5 else "FAIL"
        print(f"  PPT: {status} | {size}B | {len(prs.slides)}页 | {pics}图")
        ok = ok and status == "OK"
    else:
        print("  PPT: MISSING")
        ok = False

    if os.path.exists(html_path):
        print(f"  HTML: OK | {os.path.getsize(html_path)}B")
    else:
        print("  HTML: MISSING")
        ok = False

    if draft_result and draft_result.get("success"):
        print(f"  草稿: OK | {draft_result.get('draft_id')}")
    else:
        print("  草稿: SKIP/FAIL")

    print(f"\n  总体: {'全部通过' if ok else '有项未通过'}")
    return ok


# ======================== Image Consistency Check ========================
def check_image_consistency(base, news):
    img_dir = os.path.join(base, "images")
    print("  图片与新闻内容对应关系:")
    all_ok = True
    for i, item in enumerate(news):
        img_path = os.path.join(img_dir, f"news_{i + 1}.jpg")
        if os.path.exists(img_path):
            try:
                img = Image.open(img_path)
                w, h = img.size
                fsize = os.path.getsize(img_path)
                min_dim = min(w, h)
                status = "OK" if fsize > 10240 and min_dim >= 200 else "WARN"
                if status != "OK":
                    all_ok = False
                print(f"  [{status}] news_{i + 1}.jpg: {w}x{h} {fsize}B -> {item['title'][:30]}")
            except Exception as e:
                print(f"  [ERR] news_{i + 1}.jpg: {e}")
                all_ok = False
        else:
            print(f"  [MISS] news_{i + 1}.jpg -> {item['title'][:30]}")
            all_ok = False
    print(f"  配图一致性(尺寸): {'全部通过' if all_ok else '有项需关注'}")
    print("  [V1提醒] 尺寸校验通过不等于视觉内容匹配，运行方须用read工具逐张查看确认。")
    return all_ok


# ======================== Main ========================
def main():
    if len(sys.argv) < 2:
        print("Usage: python -X utf8 weihai_travel_runner.py news_data.json")
        sys.exit(1)

    json_path = sys.argv[1]
    base = os.path.dirname(os.path.abspath(json_path))

    with open(json_path, "r", encoding="utf-8") as f:
        data = json.load(f)

    news = data["news"]
    if "date_chinese" not in data:
        data["date_chinese"] = date_to_cn(data["date"])

    print(f"=== V1 威海旅游新闻生成 ===")
    print(f"日期: {data['date_display']} ({data['date_chinese']})")
    print(f"新闻: {len(news)}条 | 主题: {MAIN_TITLE}")

    print("\n--- 压缩配图 ---")
    img_dir = os.path.join(base, "images")
    for i in range(1, len(news) + 1):
        p = os.path.join(img_dir, f"news_{i}.jpg")
        if os.path.exists(p):
            compress_image(p)
            print(f"  news_{i}.jpg: {os.path.getsize(p)}B")

    print("\n--- 生成PPT ---")
    ppt_path = gen_ppt(base, data, news)
    print(f"  {ppt_path} ({os.path.getsize(ppt_path)}B)")

    print("\n--- 生成公众号HTML ---")
    html_path = gen_html(base, data, news)
    print(f"  {html_path} ({os.path.getsize(html_path)}B)")

    draft_result = None
    if os.name == "posix":
        print("\n--- 公众号草稿(秒哒代理) ---")
        creds = load_credentials()
        draft_result = gen_wechat_draft(base, data, news, creds)

    verify_all(base, data, news, ppt_path, html_path, draft_result)

    print("\n--- 配图一致性校验 ---")
    check_image_consistency(base, news)

    print(f"\n=== 完成 ===")


if __name__ == "__main__":
    main()